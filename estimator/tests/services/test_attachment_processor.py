from pathlib import Path
import io
import pytest
from unittest.mock import Mock, AsyncMock

from app.services.attachment_processor import extract_text_from_pdf, extract_text_from_docx, process_attachments
from docx import Document


def test_extract_text_from_pdf():
    """Should extract text content from PDF bytes."""
    pdf_path = Path(__file__).parent.parent / "fixtures" / "sample.pdf"
    assert pdf_path.exists(), "Test PDF fixture not found"

    with open(pdf_path, "rb") as f:
        pdf_bytes = f.read()

    text = extract_text_from_pdf(pdf_bytes)

    assert "Test PDF Document" in text
    assert "React" in text
    assert "PostgreSQL" in text
    assert "3 months" in text


def test_extract_text_from_docx():
    """Should extract text content from Word document bytes."""
    # Create in-memory Word doc with known content
    doc = Document()
    doc.add_heading("Project Requirements", level=1)
    doc.add_paragraph("Framework: Django")
    doc.add_paragraph("Database: MongoDB")
    doc.add_paragraph("Duration: 6 weeks")

    # Save to bytes
    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    docx_bytes = docx_buffer.getvalue()

    text = extract_text_from_docx(docx_bytes)

    assert "Project Requirements" in text
    assert "Django" in text
    assert "MongoDB" in text
    assert "6 weeks" in text


@pytest.mark.asyncio
async def test_process_attachments_multiple_files():
    """Should concatenate text from multiple files with separators."""
    from unittest.mock import patch
    from fastapi import UploadFile

    # Mock PDF UploadFile
    pdf_mock = Mock(spec=UploadFile)
    pdf_mock.filename = "spec.pdf"
    pdf_mock.content_type = "application/pdf"
    pdf_mock.read = AsyncMock(return_value=b"mock pdf bytes")

    # Mock DOCX UploadFile
    docx_mock = Mock(spec=UploadFile)
    docx_mock.filename = "requirements.docx"
    docx_mock.content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    docx_mock.read = AsyncMock(return_value=b"mock docx bytes")

    # Patch extractors to return known text
    with patch("app.services.attachment_processor.extract_text_from_pdf", return_value="PDF content here"):
        with patch("app.services.attachment_processor.extract_text_from_docx", return_value="DOCX content here"):
            result = await process_attachments([pdf_mock, docx_mock])

    assert "--- Attachment: spec.pdf ---" in result
    assert "PDF content here" in result
    assert "--- Attachment: requirements.docx ---" in result
    assert "DOCX content here" in result
